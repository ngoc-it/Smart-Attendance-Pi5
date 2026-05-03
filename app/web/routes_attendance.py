from flask import Blueprint, render_template, Response, request, session, jsonify, url_for
from datetime import datetime
import cv2
import time
import numpy as np
from app import db
from app.core import shared
from app.core.auth_utils import login_required, lecturer_required

attendance_bp = Blueprint('attendance', __name__)

@attendance_bp.route('/')
@lecturer_required
def index():
    page = request.args.get('page', 1, type=int)
    per_page = 10
    class_id = request.args.get('class_id', type=int)
    
    # Lấy danh sách lớp của GV hoặc tất cả nếu là admin
    user_role = session.get('role')
    user_id = session.get('user_id')
    if user_role == 'admin':
        classes = db.get_all_classes(per_page=100)
    else:
        classes = db.get_user_classes(user_id)
        # Bỏ phần tự động chọn lớp đầu tiên để hiện trang Hub (chọn lớp)
            
    # Dữ liệu chính
    students = []
    total_students = 0
    face_status = {}
    
    if class_id:
        students = db.get_all_students(user_id=user_id, class_id=class_id, page=page, per_page=per_page)
        total_students = db.get_student_count(user_id=user_id, class_id=class_id)
        
        # Kiểm tra trạng thái khuôn mặt
        student_ids = [s[0] for s in students]
        face_status = db.get_face_registration_status(student_ids)
    
    # Thời khóa biểu & Nghỉ phép luôn lấy để hiển thị ở Hub hoặc Detail
    schedule = db.get_user_schedule(user_id)
    leave_requests = db.get_leave_requests(user_id)
    
    pagination = {
        "current": page,
        "total_pages": (total_students + per_page - 1) // per_page,
        "total_items": total_students,
        "has_next": page * per_page < total_students,
        "has_prev": page > 1
    }
    
    return render_template('index.html', 
                           students=students, 
                           pagination=pagination,
                           total_students=total_students,
                           schedule=schedule,
                           leave_requests=leave_requests,
                           face_status=face_status,
                           classes=classes,
                           selected_class_id=class_id,
                           now_date=datetime.now().strftime("%d/%m/%Y"))

@attendance_bp.route('/leave/update', methods=['POST'])
@lecturer_required
def update_leave():
    data = request.json
    db.update_leave_status(data['request_id'], data['status'])
    return jsonify({"status": "success"})

@attendance_bp.route('/reports')
@lecturer_required
def reports():
    from flask import request, session
    date_str = request.args.get('date', datetime.now().strftime("%Y-%m-%d"))
    class_id = request.args.get('class_id', type=int)
    class_name_legacy = request.args.get('class', None)
    
    # Lọc danh sách lớp dựa trên GV
    user_role = session.get('role')
    user_id = session.get('user_id')
    
    if user_role == 'admin':
        classes = db.get_all_classes(per_page=1000)
    else:
        # Giảng viên chỉ thấy lớp mình được gán
        classes = db.get_user_classes(user_id)
    
    detailed_report = []
    stats = None
    current_class_name = ""
    
    # Ưu tiên class_id, nếu không có thử dùng class_name (tương thích cũ)
    target_class = class_id or class_name_legacy
    
    if target_class:
        # Kiểm tra quyền truy cập lớp này nếu là GV
        if user_role == 'lecturer':
            if class_id:
                allowed_ids = [c[0] for c in classes]
                if class_id not in allowed_ids:
                    return "Bạn không có quyền xem lớp này", 403
            elif class_name_legacy:
                allowed_names = [c[1] for c in classes]
                if class_name_legacy not in allowed_names:
                    return "Bạn không có quyền xem lớp này", 403

        detailed_report = db.get_detailed_attendance(date_str, target_class)
        
        # Tìm tên lớp để hiển thị
        for cls in classes:
            if (class_id and cls[0] == class_id) or (class_name_legacy and cls[1] == class_name_legacy):
                current_class_name = f"[{cls[2]}] {cls[1]}" if cls[2] else cls[1]
                if not class_id: class_id = cls[0] # Cập nhật class_id để UI đồng bộ
                break
        
        # Thống kê nhanh cho lớp được chọn
        total = len(detailed_report)
        present = len([r for r in detailed_report if r['status'] in ['Hoàn thành', 'Đi trễ', 'Có mặt', 'Chưa Checkout']])
        excused = len([r for r in detailed_report if 'Có phép' in r['status']])
        unexcused = len([r for r in detailed_report if 'Không phép' in r['status'] or r['status'] == 'Vắng'])
        
        stats = {
            "total": total,
            "present": present,
            "excused": excused,
            "unexcused": unexcused,
            "absent": total - present
        }
    
    return render_template('reports.html', 
                           stats=stats, 
                           report=detailed_report, 
                           classes=classes,
                           selected_date=date_str,
                           selected_class_id=class_id,
                           current_class_name=current_class_name)

@attendance_bp.route('/reports/export')
@login_required
def export_excel():
    from flask import send_file
    import pandas as pd
    import io
    
    date_str = request.args.get('date', datetime.now().strftime("%Y-%m-%d"))
    class_id = request.args.get('class_id', type=int)
    class_name_legacy = request.args.get('class', None)
    
    target_class = class_id or class_name_legacy
    
    if not target_class:
        return "Vui lòng chọn lớp để xuất báo cáo", 400
        
    data = db.get_detailed_attendance(date_str, target_class)
    
    # Chuyển đổi dữ liệu sang DataFrame
    df = pd.DataFrame(data)
    
    # Việt hóa tiêu đề cột
    df = df.rename(columns={
        "id": "MSSV",
        "name": "Họ và Tên",
        "in_time": "Giờ vào",
        "out_time": "Giờ ra",
        "status": "Trạng thái"
    })
    
    # Bỏ cột ảnh khi xuất excel
    if 'image' in df.columns:
        df = df.drop(columns=['image'])
        
    # Tạo file Excel trong bộ nhớ
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        df.to_excel(writer, index=False, sheet_name='Diem_Danh')
        
    output.seek(0)
    
    filename = f"DiemDanh_{target_class}_{date_str}.xlsx"
    return send_file(output, 
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                     as_attachment=True,
                     download_name=filename)

@attendance_bp.route('/reports/export_summary')
@login_required
def export_summary_excel():
    import pandas as pd
    import io
    from flask import send_file
    from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
    
    class_id = request.args.get('class_id', type=int)
    start_date = request.args.get('start_date')
    end_date = request.args.get('end_date')
    
    if not class_id:
        return "Vui lòng chọn lớp", 400
        
    class_info, students, matrix = db.get_attendance_summary_matrix(class_id, start_date, end_date)
    if not class_info:
        return "Không tìm thấy dữ liệu", 404
        
    # Chuẩn bị dữ liệu cho DataFrame
    dates = sorted(matrix.keys())
    data = []
    for idx, s in enumerate(students):
        row = {
            "STT": idx + 1,
            "MSSV": s['id'],
            "Họ và Tên": s['name']
        }
        
        absent_count = 0
        total_sessions = len(dates)
        
        for i, d in enumerate(dates):
            status = matrix[d][s['id']]
            # Header format: Buổi X (DD/MM)
            date_obj = datetime.strptime(d, "%Y-%m-%d")
            header = f"Buổi {i+1}\n({date_obj.strftime('%d/%m')})"
            row[header] = status or "K" # Mặc định là vắng nếu không có log
            if status in ["K", "P", ""]:
                absent_count += 1
                
        row["Tổng buổi"] = total_sessions
        row["Vắng"] = absent_count
        row["% Nghỉ"] = round((absent_count / total_sessions * 100), 1) if total_sessions > 0 else 0
        data.append(row)
        
    df = pd.DataFrame(data)
    
    # Xuất ra Excel với định dạng chuyên nghiệp
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        df.to_excel(writer, index=False, sheet_name='Tong_Hop_Diem_Danh', startrow=4)
        
        workbook = writer.book
        worksheet = writer.sheets['Tong_Hop_Diem_Danh']
        
        # 1. Thêm Header thông tin lớp
        title_font = Font(size=14, bold=True)
        worksheet['A1'] = "BÁO CÁO TỔNG HỢP ĐIỂM DANH"
        worksheet['A1'].font = title_font
        worksheet['A2'] = f"Học phần: {class_info['course_name']} ({class_info['class_code']})"
        worksheet['A3'] = f"Lớp: {class_info['class_name']} | Dữ liệu tính đến: {datetime.now().strftime('%d/%m/%Y')}"
        
        # 2. Định dạng Header Table
        header_fill = PatternFill(start_color="2D3748", end_color="2D3748", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True)
        center_align = Alignment(horizontal='center', vertical='center', wrap_text=True)
        thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
        
        for cell in worksheet[5]: # Dòng 5 là header (index 0-based is 4, but pandas startrow=4 makes it 5)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center_align
            cell.border = thin_border
            
        # 3. Định dạng dữ liệu & Conditional Formatting (giả lập)
        present_fill = PatternFill(start_color="C6FFD8", end_color="C6FFD8", fill_type="solid") # Xanh lá nổi bật
        late_fill = PatternFill(start_color="FEFCBF", end_color="FEFCBF", fill_type="solid") # Vàng nhạt
        excused_fill = PatternFill(start_color="FFFBEB", end_color="FFFBEB", fill_type="solid") # Cam nhạt
        absent_fill = PatternFill(start_color="FED7D7", end_color="FED7D7", fill_type="solid") # Đỏ nhạt
        warning_font = Font(color="C53030", bold=True)
        
        # Duyệt qua các dòng dữ liệu
        for row_idx, row_cells in enumerate(worksheet.iter_rows(min_row=6, max_row=worksheet.max_row)):
            for cell_idx, cell in enumerate(row_cells):
                cell.border = thin_border
                val = str(cell.value)
                
                # Căn giữa cho các cột trạng thái và tổng kết
                if cell_idx >= 3:
                    cell.alignment = center_align
                
                # Tô màu trạng thái
                if val == 'x': cell.fill = present_fill
                elif val == 'T': cell.fill = late_fill
                elif val == 'P': cell.fill = excused_fill
                elif val == 'K': cell.fill = absent_fill
                
                # Cảnh báo nghỉ quá 20% (Cột cuối cùng)
                if cell_idx == len(row_cells) - 1:
                    try:
                        if float(val) >= 20:
                            cell.font = warning_font
                            cell.fill = absent_fill
                    except: pass
        
        # 4. Điều chỉnh độ rộng cột
        from openpyxl.utils import get_column_letter
        worksheet.column_dimensions['B'].width = 15 # MSSV
        worksheet.column_dimensions['C'].width = 25 # Họ Tên
        
        # Tự động điều chỉnh độ rộng cho các cột Buổi học và Tổng kết
        for col_idx in range(4, worksheet.max_column + 1):
            col_letter = get_column_letter(col_idx)
            worksheet.column_dimensions[col_letter].width = 12
            
        # 5. Cố định Pane (STT, MSSV, Họ Tên)
        worksheet.freeze_panes = 'D6'

    output.seek(0)
    filename = f"TongHop_{class_info['class_code']}_{datetime.now().strftime('%Y%m%d')}.xlsx"
    return send_file(output, 
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                     as_attachment=True,
                     download_name=filename)

@attendance_bp.route('/video_feed')
def video_feed():
    def generate():
        while True:
            time.sleep(0.04) # 25 FPS
            img = None
            
            # TRUY XUẤT TỪ SHARED
            with shared.frame_lock:
                if shared.last_frame is not None:
                    img = shared.last_frame.copy()
            
            if img is not None:
                # LẤY KẾT QUẢ AI TỪ SHARED
                results = shared.last_results
                for res in results:
                    x1, y1, x2, y2 = res['box']
                    name = res['id']
                    cv2.rectangle(img, (x1, y1), (x2, y2), (0, 255, 0), 2)
                    cv2.putText(img, name, (x1, y1 - 10), 
                                cv2.FONT_HERSHEY_SIMPLEX, 0.6, (0, 255, 0), 2)
            else:
                # Màn hình chờ nếu camera chưa có data trong shared
                img = np.zeros((480, 640, 3), np.uint8)
                img[:] = (40, 40, 40)
                cv2.putText(img, "CONNECTING VIA SHARED MEMORY...", (100, 240), 
                            cv2.FONT_HERSHEY_SIMPLEX, 0.7, (255, 255, 255), 2)

            ret, buffer = cv2.imencode('.jpg', img)
            if ret:
                yield (b'--frame\r\n'
                       b'Content-Type: image/jpeg\r\n\r\n' + buffer.tobytes() + b'\r\n')
    
    return Response(generate(), mimetype='multipart/x-mixed-replace; boundary=frame')

@attendance_bp.route('/api/student/<student_id>', methods=['DELETE'])
def delete_student(student_id):
    try:
        from app import ai_engine
        import os
        from config import Config
        
        # 1. Delete from DB
        db.delete_student(student_id)
        
        # 2. Delete embedding file
        emb_path = os.path.join(Config.EMBEDDINGS_DIR, f"{student_id}.npy")
        if os.path.exists(emb_path):
            os.remove(emb_path)
            
        # 3. Reload AI engine database
        ai_engine.load_database()
        
        return {"status": "success", "message": "Deleted successfully"}
    except Exception as e:
        return {"status": "error", "message": str(e)}, 500

@attendance_bp.route('/api/student/<student_id>', methods=['PUT'])
def update_student(student_id):
    from flask import request
    try:
        data = request.json
        name = data.get('name')
        dob = data.get('dob')
        student_class = data.get('student_class')
        
        db.update_student_info(student_id, name, dob, student_class)
        
        return {"status": "success", "message": "Updated successfully"}
    except Exception as e:
        return {"status": "error", "message": str(e)}, 500

@attendance_bp.route('/attendance/action', methods=['POST'])
def manual_action():
    from flask import request
    try:
        data = request.json
        student_id = data.get('id')
        action = data.get('action')
        
        # Lấy tên sinh viên từ DB
        name = db.get_student_name(student_id)
                
        db.update_status(student_id, name, action)
        return {"status": "success"}
    except Exception as e:
        return {"status": "error", "message": str(e)}, 500

@attendance_bp.route('/images/<path:filename>')
@login_required
def serve_image(filename):
    from flask import send_from_directory
    from config import Config
    return send_from_directory(Config.LOG_IMAGES_DIR, filename)

@attendance_bp.route('/api/attendance/manual', methods=['POST'])
@lecturer_required
def attendance_manual():
    from flask import request
    data = request.json
    student_id = data.get('id')
    action = data.get('action')
    date_str = data.get('date', datetime.now().strftime("%Y-%m-%d"))
    reason = data.get('reason')
    
    success = db.update_attendance_manual(student_id, action, date_str, reason)
    if success:
        return jsonify({"status": "success"})
    return jsonify({"status": "error", "message": "Lỗi cập nhật"}), 500

@attendance_bp.route('/api/attendance/sessions/<int:class_id>')
@lecturer_required
def get_sessions(class_id):
    sessions = db.get_editable_sessions(class_id)
    return jsonify(sessions)
